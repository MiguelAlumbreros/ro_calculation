"""Entry points for running the calculation.

Two ways in:

- `run_asof(asof, df_power, df_gas, df_eua, target_quarter)` returns a result dict and
  keeps everything in memory. Either `asof` or `target_quarter`, or both, may be a list,
  which fans out and returns a dict keyed by whichever was a list.
- `run(target_quarter)` loads the price CSV itself and writes an output workbook.

Load the price frames once with `load_price_frames()` and pass them to `run_asof` as
often as you like.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl import load_workbook

from .config import (
    COMMODITY_LEGS,
    data_dir,
    output_dir,
    pfc_path,
    templates_dir,
)
from .constants import constants_for, ivpee_for_quarter
from .decomposition import (
    COMPONENT_COLUMNS,
    averaging_window,
    gather_price_components,
    normalize_prices,
    parse_target_quarter,
)
from .formula import calculate_ro
from .pfc import load_pfc
from .realised import realised_ro
from .trading_calendar import asof_status, trading_days_remaining


__all__ = [
    "ivpee_for_quarter",
    "compute_ro_scenarios",
    "write_ro_workbook",
    "load_price_frames",
    "run_asof",
    "run",
]


# Where each (commodity, strip) lands in data/ro_output_template.xlsx. Positions
# are fixed, so edit this map and the template together. Columns are 1=quarter, 4=fixed,
# 5=float, 6=reference, 7=pct_open.
_TEMPLATE_ROW_MAP: dict[tuple[str, str], int] = {
    ("power", "Y"): 2, ("power", "Q"): 3, ("power", "M1"): 4,
    ("power", "M2"): 5, ("power", "M3"): 6, ("power", "total"): 7,
    ("gas", "Y"): 8, ("gas", "Q"): 9, ("gas", "M1"): 10,
    ("gas", "M2"): 11, ("gas", "M3"): 12, ("gas", "total"): 13,
    ("EUA", "M3"): 14, ("EUA", "total"): 15,
}
_RO_ROW = 16

# Display label overrides for legs in run_asof's output dict: (commodity, leg) → label
_LEG_LABEL: dict[tuple[str, str], str] = {
    ("gas", "Y"): "Cal",
}


def _constants_for(target_quarter: str, constants: Mapping[str, float] | None) -> Mapping[str, float]:
    if constants is not None:
        return constants
    return constants_for(target_quarter)


def compute_ro_scenarios(
    components: pd.DataFrame,
    constants: Mapping[str, float],
) -> dict[str, float]:
    """Run calculate_ro three times (fixed / float / reference) on the per-commodity totals.

    The three scenarios for the same target quarter:

    - `fixed`     — RO implied by the closes already observed. NaN until at least one
                    full leg has elapsed.
    - `float`     — RO if today's forward prices held for the rest of every window. NaN
                    once every window has closed.
    - `reference` — the day-weighted blend of the two, and the figure to quote.

    Each is evaluated on its own, so a valid `reference` can sit beside a NaN `fixed`. A
    NaN in any of P_m/P_pvb/P_co2 makes that scenario NaN.
    """
    # First row per commodity.
    totals: dict[str, pd.Series] = {}
    for _, row in components[components["strip"] == "total"].iterrows():
        totals.setdefault(row["commodity"], row)

    def total(commodity: str, col: str) -> float:
        row = totals.get(commodity)
        if row is None:
            raise KeyError(f"Missing total row for {commodity}.")
        return float(row[col])

    scenarios: dict[str, float] = {}
    for label in ("fixed", "float", "reference"):
        p_m = total("power", label)
        p_pvb = total("gas", label)
        p_co2 = total("EUA", label)
        if pd.isna(p_m) or pd.isna(p_pvb) or pd.isna(p_co2):
            scenarios[label] = float("nan")
        else:
            scenarios[label] = float(calculate_ro(
                {"P_m": p_m, "P_pvb": p_pvb, "P_co2": p_co2},
                constants,
            ))
    return scenarios


def load_price_frames(csv_path: Path | None = None) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Load the price CSV and split it into the power / gas / EUA frames.

    Reads `data/futures.csv` unless `csv_path` says otherwise. The frames come back
    normalised (`decomposition.normalize_prices`): narrowed to the cal, quarter and
    month strips, date columns parsed, filter columns held as categories.

    Call this once and reuse the three frames. Tidying them afterwards, for instance
    deduplicating EUA closes, keeps them normalised.
    """
    csv_path = csv_path or data_dir() / "futures.csv"
    df = pd.read_csv(csv_path, low_memory=False)
    return (
        normalize_prices(df[df["commodity"] == "power"]),
        normalize_prices(df[df["commodity"] == "gas"]),
        normalize_prices(df[df["commodity"] == "EUA"]),
    )


def write_ro_workbook(
    components: pd.DataFrame,
    ro_scenarios: dict[str, float],
    target_quarter: str,
    template_path: Path,
    output_path: Path,
) -> None:
    """Fill the template with components and the three RO scenarios.

    Cell formatting in the template is preserved. The workbook is written to
    `output_path`, creating its parent directory if needed.
    """
    wb = load_workbook(template_path)
    ws = wb.active

    for _, r in components.iterrows():
        rownum = _TEMPLATE_ROW_MAP.get((r["commodity"], r["strip"]))
        if rownum is None:
            continue
        ws.cell(row=rownum, column=1, value=target_quarter)
        for col_idx, col_name in ((4, "fixed"), (5, "float"), (6, "reference"), (7, "pct_open")):
            v = r[col_name]
            ws.cell(row=rownum, column=col_idx, value=None if pd.isna(v) else float(v))

    ws.cell(row=_RO_ROW, column=1, value=target_quarter)
    for col_idx, label in ((4, "fixed"), (5, "float"), (6, "reference")):
        v = ro_scenarios.get(label)
        ws.cell(row=_RO_ROW, column=col_idx, value=None if v is None or pd.isna(v) else float(v))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _realised_ro_row(target_quarter: str, ro: float) -> pd.DataFrame:
    """A one-row components frame carrying a published RO and nothing else.

    It has the same columns as a calculated result, so frames from both paths can be
    concatenated.
    """
    return pd.DataFrame([{
        "quarter": target_quarter,
        "commodity": "RO",
        "strip": "total",
        "fixed": ro,
        "float": ro,
        "reference": ro,
        "pct_open": 0.0,
        "price_source": "realised",
    }])


def _run_asof_single(
    asof: date,
    df_power: pd.DataFrame,
    df_gas: pd.DataFrame,
    df_eua: pd.DataFrame,
    target_quarter: str,
    constants: Mapping[str, float] | None = None,
    pfc_df: pd.DataFrame | None = None,
) -> dict:
    """Decompose prices as of `asof`, compute RO scenarios, and report leg status.

    A quarter already realised as of `asof` returns its published RO and skips the
    decomposition entirely (see `realised.realised_ro`).

    `pfc_df` is the loaded hourly Price Forward Curve (see `pfc.load_pfc`); when given,
    it supplies the floating portion of every leg it has a column for. See
    `gather_price_components`.

    Returns a dict with:
      run_as_of         the `asof` echoed back
      run_as_of_status  'trading_day' | 'weekend' | 'holiday'
      realised          True when the RO below is the published one, not an estimate
      ro_scenarios      {'fixed', 'float', 'reference'} in EUR/MWhE; all three equal
                        when realised
      legs              per '<commodity>_<leg>': price_source + trading_days_remaining;
                        empty when realised
      df                the components frame with an appended 'RO' row; the 'RO' row
                        alone when realised
    """
    published = realised_ro(target_quarter, asof)
    if published is not None:
        return {
            "run_as_of": asof,
            "run_as_of_status": asof_status(asof),
            "realised": True,
            "ro_scenarios": {"fixed": published, "float": published, "reference": published},
            "legs": {},
            "df": _realised_ro_row(target_quarter, published),
        }

    constants = _constants_for(target_quarter, constants)

    components = gather_price_components(df_power, df_gas, df_eua, target_quarter, asof, pfc_df)
    scenarios = compute_ro_scenarios(components, constants)

    commodity_totals = components.loc[components["strip"] == "total"]
    # Unweighted mean across power/gas/EUA. Read it as a rough "how much of this RO can
    # still move", not as a sensitivity: the commodities do not enter the formula equally.
    ro_pct_open = float(commodity_totals["pct_open"].mean())
    # Flag the RO row when any commodity's price did not come straight from its primary
    # exchange, so a fallback or propagated leg shows up at the top level.
    non_primary = commodity_totals.loc[~commodity_totals["primary_direct"]]
    if non_primary.empty:
        ro_source = "consistent"
    else:
        detail = "; ".join(f"{r.commodity}: {r.price_source}" for r in non_primary.itertuples())
        ro_source = f"mixed ({detail})"
    ro_row = pd.DataFrame([{
        "quarter": target_quarter,
        "commodity": "RO",
        "strip": "total",
        "fixed": scenarios["fixed"],
        "float": scenarios["float"],
        "reference": scenarios["reference"],
        "pct_open": ro_pct_open,
        "price_source": ro_source,
    }])
    full_df = pd.concat([components, ro_row], ignore_index=True)

    legs_info = {}
    for commodity, legs in COMMODITY_LEGS.items():
        for leg in legs:
            _, win_end = averaging_window(target_quarter, commodity, leg)
            strip = leg if len(legs) > 1 else "total"
            comp_row = components[
                (components["commodity"] == commodity) & (components["strip"] == strip)
            ]
            price_source = comp_row.iloc[0]["price_source"] if not comp_row.empty else None
            label = _LEG_LABEL.get((commodity, leg), leg)
            legs_info[f"{commodity}_{label}"] = {
                "price_source": price_source,
                "trading_days_remaining": trading_days_remaining(asof, win_end, commodity),
            }

    return {
        "run_as_of": asof,
        "run_as_of_status": asof_status(asof),
        "realised": False,
        "ro_scenarios": scenarios,
        "legs": legs_info,
        "df": full_df,
    }


def run_asof(
    asof: date | list[date],
    df_power: pd.DataFrame,
    df_gas: pd.DataFrame,
    df_eua: pd.DataFrame,
    target_quarter: str | list[str],
    constants: Mapping[str, float] | None = None,
    use_pfc: bool = True,
    pfc_csv_path: Path | None = None,
) -> dict:
    """The main entry point: what will the RO for `target_quarter` be, seen from `asof`?

    `asof` sets how much of each averaging window has already elapsed, and so how much of
    the answer is locked in. See `compute_ro_scenarios` for what the three RO figures
    mean, and `_run_asof_single` for the full result shape.

    `constants` defaults to `constants.constants_for(target_quarter)`: the hardcoded
    tolls and levies, the valores propios from `ro_parameters.xlsx`, and the IVPEE
    rate from `IVPEE.xlsx` (both shipped in the package). Pass a mapping to override it; that mapping is then used
    as given, including its IVPEE.

    `use_pfc` takes the floating (still-open) portion of every leg from the hourly Price
    Forward Curve; the fixed portion still comes from futures. It is on by default. Pass
    False to hold the last observed close flat instead. `pfc_csv_path` overrides the
    default PFC location (`config.pfc_path()`) and is ignored when `use_pfc` is False.

    Either argument also accepts a list, which fans out:

    - Both scalar: returns the single `_run_asof_single` result dict.
    - `asof` a list, `target_quarter` scalar: returns a dict keyed by asof.
    - `target_quarter` a list, `asof` scalar: returns a dict keyed by target_quarter.
    - Both lists: returns a dict keyed by asof, each value a dict keyed by target_quarter.
    """
    pfc_df = load_pfc(pfc_csv_path or pfc_path()) if use_pfc else None

    # Normalise the three frames up front. Frames from `load_price_frames` are already
    # normalised and pass through unchanged.
    df_power = normalize_prices(df_power)
    df_gas = normalize_prices(df_gas)
    df_eua = normalize_prices(df_eua)

    asof_is_list = isinstance(asof, list)
    quarter_is_list = isinstance(target_quarter, list)

    if not asof_is_list and not quarter_is_list:
        return _run_asof_single(asof, df_power, df_gas, df_eua, target_quarter, constants, pfc_df)

    asofs = asof if asof_is_list else [asof]
    quarters = target_quarter if quarter_is_list else [target_quarter]

    results = {
        a: {q: _run_asof_single(a, df_power, df_gas, df_eua, q, constants, pfc_df) for q in quarters}
        for a in asofs
    }

    if not asof_is_list:
        return results[asofs[0]]
    if not quarter_is_list:
        return {a: results[a][quarters[0]] for a in asofs}
    return results


def run(
    target_quarter: str,
    asof: date | None = None,
    csv_path: Path | None = None,
    template_path: Path | None = None,
    out_dir: Path | None = None,
    constants: Mapping[str, float] | None = None,
    use_pfc: bool = True,
    pfc_csv_path: Path | None = None,
) -> pd.DataFrame:
    """End-to-end: load CSV, decompose, compute RO scenarios, write the workbook.

    The workbook lands in `out_dir` as `ro_<quarter>_asof_<YYYYMMDD>.xlsx`, and the
    components frame is also returned. `asof` defaults to today.

    A quarter already realised as of `asof` skips the decomposition: the workbook carries
    the published RO and no price components. See `run_asof` for `use_pfc` and
    `pfc_csv_path`.
    """
    if asof is None:
        asof = date.today()

    template_path = template_path or templates_dir() / "ro_output_template.xlsx"
    out_dir = out_dir or output_dir()

    published = realised_ro(target_quarter, asof)
    if published is not None:
        components = pd.DataFrame(columns=COMPONENT_COLUMNS)
        scenarios = {"fixed": published, "float": published, "reference": published}
    else:
        constants = _constants_for(target_quarter, constants)
        df_power, df_gas, df_eua = load_price_frames(csv_path)
        pfc_df = load_pfc(pfc_csv_path or pfc_path()) if use_pfc else None
        components = gather_price_components(df_power, df_gas, df_eua, target_quarter, asof, pfc_df)
        scenarios = compute_ro_scenarios(components, constants)

    output_path = out_dir / f"ro_{target_quarter}_asof_{asof.strftime('%Y%m%d')}.xlsx"
    write_ro_workbook(components, scenarios, target_quarter, template_path, output_path)

    ro_row = pd.DataFrame([{
        "quarter": target_quarter,
        "commodity": "RO",
        "strip": "total",
        "fixed": scenarios["fixed"],
        "float": scenarios["float"],
        "reference": scenarios["reference"],
        "pct_open": None,
    }])
    return pd.concat([components, ro_row], ignore_index=True)
